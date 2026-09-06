"""Muavin döküm alımı + denetimi için offline doğrulama.

Kapsam:
 1) 'Yeni format' hesap bloklu Luca defterinin (TARİH/FİŞ NO/SR/AÇIKLAMA/
    BORÇ TUT./ALACAK TUT. + FT.MIZ NO satırları) cetvel_dosya_parse ile
    tanındığını doğrular (191/391 parser'ı, genel otomatik tanıma yedeği).
 2) Tek hücreli boş döküm (muavin stübü) için cetvel_dosya_parse'in kayıt
    üretmediğini ve _muavin_dosya_denetle'nin False döndüğünü doğrular.
 3) Tarih sütunu datetime objesi olan (Excel seri tarih) dolu dökümde
    hareket algılamasını ve 191/192 hesap denetimini doğrular.
 4) _RE_TARIH_ISARETI sayacının gerçek defter metinlerini ayırt ettiğini
    doğrular (rapor verisi bekleme sinyali).
"""
import os
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

YOL = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, YOL)

import openpyxl
from datetime import datetime

from dosya import cetvel_dosya_parse, _kullanilabilir_mi
from luca_cekme import _muavin_dosya_denetle, _RE_TARIH_ISARETI

BASARILI = True


def kontrol(ad, kosul, detay=""):
    global BASARILI
    durum = "TAMAM" if kosul else "HATA"
    if not kosul:
        BASARILI = False
    print(f"  [{durum}] {ad} {detay}")


def yeni_format_defter(yol):
    """Tipik Luca raporMizanDetayHazirla defter düzeni: 'Hesap Kodu 191-xx'
    blogu + TARİH/FİŞ NO/SR/AÇIKLAMA/BORÇ TUT./ALACAK TUT. basliklari."""
    wb = openpyxl.Workbook()
    ws = wb.active
    baslik = ["TARİH", "FİŞ NO", "SR", "AÇIKLAMA", "BORÇ TUT.",
              "ALACAK TUT.", "REFERANS KODU", "REFERANS İSMİ", "İŞLEM TİPİ"]
    satirlar = [
        ["MUAVİN DEFTER", None, None, None, None, None, None, None, None],
        ["FİRMA ÜNVANI : TEST MÜKELLEF LTD.", None, None, None, None, None,
         None, None, None],
        ["DÖNEM 01.08.2026 31.08.2026", None, None, None, None, None, None,
         None, None],
        ["TARİH 06.09.2026", None, None, None, None, None, None, None, None],
        ["191-01-001 İNDİRİLECEK KDV", None, None, None, None, None, None,
         None, None],
        baslik,
        ["01.08.2026", "FT1", "1",
         "EK DİJİTAL MALZEME FT.MIZ NO:GFE202600000011 KDVSI",
         1200.00, "", "", "", ""],
        ["05.08.2026", "FT2", "2",
         "OFİS MALZEMESİ FT.MIZ NO:GFE202600000012 KDVSI",
         800.00, "", "", "", ""],
    ]
    for satir in satirlar:
        ws.append([s if s is not None else "" for s in satir])
    wb.save(yol)
    wb.close()


def stub_defter(yol):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["MUAVİN DEFTER"])
    wb.save(yol)
    wb.close()


def datetime_defter(yol):
    wb = openpyxl.Workbook()
    ws = wb.active
    satirlar = [
        ["MUAVİN DEFTER", None, None, None, ],
        ["FİRMA ÜNVANI : TEST", None, None, None],
        ["DÖNEM 01.08.2026 31.08.2026", None, None, None],
        ["TARİH 06.09.2026", None, None, None],
        ["191-01-003 İNDİRİLECEK KDV", None, None, None],
        ["TARİH", "FİŞ NO", "AÇIKLAMA", "BORÇ TUT."],
        [datetime(2026, 8, 3), "FT7", "DIZAN PTM FT.MIZ NO:GFE202600000031",
         450.00],
        [datetime(2026, 8, 9), "FT8", "MASA FT.MIZ NO:GFE202600000032",
         250.00],
    ]
    for satir in satirlar:
        ws.append([s if s is not None else "" for s in satir])
    wb.save(yol)
    wb.close()


def reporter(depo):
    def b(msj):
        depo.append(msj)
    return b


if __name__ == "__main__":
    gecici = os.path.join(os.environ.get("TEMP", YOL), "kdv_muavin_birim")
    os.makedirs(gecici, exist_ok=True)

    print("== 1) YENİ FORMAT LUCA DEFTERİ ==")
    yol1 = os.path.join(gecici, "luca_muavin_191_202608.xlsx")
    yeni_format_defter(yol1)
    sonuc = cetvel_dosya_parse(yol1)
    kayitlar = sonuc.get("kayitlar", [])
    for k in kayitlar:
        print(f"  belge={k['belge_no']} tarih={k['tarih']} "
              f"kdv={k['kdv']} unvan={k['unvan'][:30]}")
    kontrol("yeni format kayit sayisi", len(kayitlar) == 2,
            f"-> {len(kayitlar)}")
    kontrol("yeni format ilk belge",
            kayitlar and kayitlar[0]["belge_no"] == "GFE202600000011")
    kontrol("yeni format ilk kdv",
            kayitlar and abs(float(kayitlar[0]["kdv"]) - 1200.00) < 0.01,
            f"-> {kayitlar[0]['kdv'] if kayitlar else None}")
    kontrol("kullanilabilir sayilir", _kullanilabilir_mi(sonuc))

    print("\n== 2) TEK HÜCRELİ STÜB ==")
    yol2 = os.path.join(gecici, "luca_muavin_391_202608.xlsx")
    stub_defter(yol2)
    st_sonuc = cetvel_dosya_parse(yol2)
    kontrol("stüb kayit üretmez", len(st_sonuc.get("kayitlar", [])) == 0,
            f"-> {len(st_sonuc.get('kayitlar', []))}")
    kontrol("stüb kullanilabilir değil",
            not _kullanilabilir_mi(st_sonuc))
    notlar = []
    ok = _muavin_dosya_denetle(yol2, "391", reporter(notlar))
    kontrol("denetle stüb False", ok is False)
    kontrol("denetle stüb uyarısı",
            any("boş görünüyor" in n for n in notlar),
            f"-> {notlar[:1]}")

    print("\n== 3) DATETIME SÜTUNLU DOLU DÖKÜM ==")
    yol3 = os.path.join(gecici, "luca_muavin_191_dt_202608.xlsx")
    datetime_defter(yol3)
    notlar3 = []
    ok3 = _muavin_dosya_denetle(yol3, "191", reporter(notlar3))
    kontrol("datetime döküm True", ok3 is True)
    kontrol("191 hesabı tanındı",
            any("doğrulandı" in n for n in notlar3), f"-> {notlar3}")
    kontrol("192 yok Not seviyesinde",
            any("Hesap 192 dökümde yok" in n for n in notlar3))

    print("\n== 4) TARİH İŞARETİ SAYACI ==")
    dolu_metin = ("01.08.2026 Acıklama\n05.08.2026 Açıklama\n"
                  "10.08.2026 Açıklama")
    kontrol("3 gelir satırı sayılır",
            len(_RE_TARIH_ISARETI.findall(dolu_metin)) == 3)
    bos_metin = "MUAVİN DEFTER\nDönem 01.08.2026 31.08.2026"
    kontrol("boş ekranda 3 altı kalır",
            len(_RE_TARIH_ISARETI.findall(bos_metin)) < 3,
            f"-> {len(_RE_TARIH_ISARETI.findall(bos_metin))}")
    kontrol("IZO tarih sayılmaz",
            len(_RE_TARIH_ISARETI.findall("2026-08-01 2026-08-05")) == 0)

    print()
    print("SONUÇ:", "TAMAM" if BASARILI else "HATA VAR")
    sys.exit(0 if BASARILI else 1)